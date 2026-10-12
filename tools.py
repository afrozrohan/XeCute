import json
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from langchain_core.tools import tool

BASE_DIR = Path(__file__).resolve().parent
WORKSPACE = (BASE_DIR / "workspace").resolve()
WORKSPACE.mkdir(parents=True, exist_ok=True)

MAX_TEXT_CHARS = 15000
MAX_EXCEL_ROWS = 200
MAX_LISTED_FILES = 200


def safe_path(filename: str) -> Path:
    """Resolve a filename to a path inside the workspace."""
    path = Path(filename)

    if path.is_absolute():
        resolved = path.resolve()
    else:

        if path.parts and path.parts[0].lower() == "workspace":
            path = Path(*path.parts[1:])
        resolved = (WORKSPACE / path).resolve()

    if not resolved.is_relative_to(WORKSPACE):
        raise ValueError("Path must be inside the workspace.")

    return resolved


def _truncate(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n...[truncated, {len(text) - limit} more characters]"


@tool
def list_workspace_files(folder: str = ".") -> list[str]:
    """List files inside the workspace (paths are relative to the workspace)."""
    path = safe_path(folder)

    if not path.exists():
        return []

    files = [
        file.relative_to(WORKSPACE).as_posix()
        for file in path.rglob("*")
        if file.is_file()
    ]
    return files[:MAX_LISTED_FILES]


@tool
def read_file(filename: str) -> str:
    """Read a plain text file from the workspace. Not for PDF or Excel files."""
    path = safe_path(filename)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {filename}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raise ValueError("This is a PDF. Use read_pdf instead.")
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        raise ValueError("This is an Excel file. Use read_excel instead.")

    try:
        return _truncate(path.read_text(encoding="utf-8"))
    except UnicodeDecodeError:
        raise ValueError(f"{filename} is not a UTF-8 text file.")


@tool
def write_file(filename: str, content: str) -> str:
    """Write a text file inside the workspace."""
    path = safe_path(filename)

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")

    return f"File written: {path.relative_to(WORKSPACE).as_posix()}"


@tool
def read_pdf(filename: str) -> str:
    """Extract text from a PDF inside the workspace."""
    path = safe_path(filename)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {filename}")
    if path.suffix.lower() != ".pdf":
        raise ValueError("read_pdf only supports .pdf files.")

    parts = []
    with pymupdf.open(path) as document:
        for number, page in enumerate(document, start=1):
            parts.append(f"--- Page {number} ---\n{page.get_text()}")

    text = "\n".join(parts)

    if len(text.replace("-", "").strip()) < 20 or not any(
        ch.isalnum() for ch in text
    ):
        return (
            "[No extractable text found. The PDF may be a scanned image "
            "and would need OCR.]"
        )

    return _truncate(text)


@tool
def read_excel(filename: str, sheet_name: str = "") -> str:
    """Read an Excel workbook (JSON rows). Optionally give a sheet name."""
    path = safe_path(filename)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {filename}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Sheet {sheet_name!r} not found. "
                    f"Available: {workbook.sheetnames}"
                )
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active

        rows = []
        total = 0
        for row in sheet.iter_rows(values_only=True):
            total += 1
            if len(rows) < MAX_EXCEL_ROWS:
                rows.append(list(row))

        result = {
            "sheet": sheet.title,
            "sheets_available": workbook.sheetnames,
            "total_rows": total,
            "rows_returned": len(rows),
            "rows": rows,
        }
    finally:
        workbook.close()

    # default=str handles dates, decimals and other non-JSON values.
    return _truncate(json.dumps(result, ensure_ascii=False, default=str))


@tool
def write_excel(filename: str, row: list) -> str:
    """Append a row to the active worksheet of an existing workbook."""
    path = safe_path(filename)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {filename}")

    workbook = load_workbook(path)
    try:
        workbook.active.append(row)
        workbook.save(path)
    finally:
        workbook.close()

    return f"Added row to {filename}"
